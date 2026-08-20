"""
Generate cover letter DOCX (matches resume styling).
Run: python scripts/build-cover-letter.py
"""

from pathlib import Path
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "daphne-chepkirui-cover-letter.docx"

BLUE = RGBColor(37, 99, 235)
BLACK = RGBColor(15, 23, 42)
PORTFOLIO = "https://portfolio-omega-umber-28.vercel.app/"


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = BLACK

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = name.add_run("Daphne Chepkirui")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = BLUE

    private = {}
    private_path = DOCS / "resume-private.json"
    if private_path.exists():
        private = json.loads(private_path.read_text(encoding="utf-8"))
    phone_display = private.get("phone_display") or private.get("phone") or ""

    contact = doc.add_paragraph()
    contact.add_run("Nairobi, Kenya\n")
    p = doc.add_paragraph()
    add_hyperlink(p, "chepkiruidaphne91@gmail.com", "mailto:chepkiruidaphne91@gmail.com")
    if phone_display:
        p.add_run(f" | {phone_display}")
    p.add_run("\nPortfolio: ")
    add_hyperlink(p, "portfolio-omega-umber-28.vercel.app", PORTFOLIO)
    p.add_run(" | ")
    add_hyperlink(
        p,
        "LinkedIn",
        "https://www.linkedin.com/in/daphne-chepkirui-382178313",
    )

    doc.add_paragraph()

    doc.add_paragraph("[Date]")
    doc.add_paragraph()
    doc.add_paragraph("[Hiring Manager Name]")
    doc.add_paragraph("[Company Name]")
    doc.add_paragraph("[Company Address or Remote]")
    doc.add_paragraph()

    doc.add_paragraph("Dear [Hiring Manager Name / Hiring Team],")
    doc.add_paragraph()

    paragraphs = [
        (
            "I am writing to apply for the [Job Title] role at [Company Name]. "
            "I am a JavaScript and TypeScript developer based in Nairobi. I ship products "
            "across the stack: Next.js on the web, NestJS for APIs, and Expo React Native on mobile. "
            "Your team's focus on [one specific thing from the job post] lines up with the work I "
            "have been doing, and I would like to bring that to [Company Name]."
        ),
        (
            "At Dynamic Mobility Technology I have shipped real products. KISRS is live in UAT for "
            "hospitals and labs (referral dashboards and RBAC). I built BeautiLink, a beauty and wellness "
            "marketplace: Next.js admin, NestJS APIs, and an Expo app for shop, chat, and dashboards. "
            "I also delivered insurance admin and agents portals, e-Sahal wallet flows, and CarbonFlow reporting. "
            "I am used to dense ops UIs, RBAC, and NDA work."
        ),
        (
            "Before that I shipped offline-first POS and MRP at Tenzi Limited, a CBC student portal for Upeo, "
            "and several public Next.js apps. I care about how products feel and how they hold up: accessibility, "
            "code splitting, lazy loading, and clear handoffs are part of how I work."
        ),
    ]

    for text in paragraphs:
        doc.add_paragraph(text)

    closing = doc.add_paragraph()
    closing.add_run(
        "I would welcome a conversation about how I can help [Company Name] with "
        "[specific team goal from the posting]. My portfolio at "
    )
    add_hyperlink(closing, "portfolio-omega-umber-28.vercel.app", PORTFOLIO)
    closing.add_run(
        " includes project breakdowns and UI previews, and I am happy to walk through "
        "enterprise work in a live demo if that would be useful."
    )

    doc.add_paragraph()
    doc.add_paragraph("Thank you for your time and consideration.")
    doc.add_paragraph()
    doc.add_paragraph("Sincerely,")
    sig = doc.add_paragraph()
    r = sig.add_run("Daphne Chepkirui")
    r.bold = True

    DOCS.mkdir(exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
