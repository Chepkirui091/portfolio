"""
Generate ATS-friendly resume (PDF + DOCX) from portfolio-aligned content.
Run: python scripts/build-resume.py
"""

from __future__ import annotations

import json
import os
import re
from pathlib import Path
import shutil
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
RESUME_PDF = DOCS / "daphne-chepkirui-resume.pdf"
PUBLIC_RESUME = ROOT / "public" / "Daphne_Chepkirui_Resume.pdf"

# Titles and links only (skills body stays black)
BLUE = (37, 99, 235)  # #2563eb
BLUE_HEX = "#2563eb"
BLACK = (15, 23, 42)  # #0f172a
BLACK_HEX = "#0f172a"

CONTACT = {
    "name": "DAPHNE CHEPKIRUI",
    "title": "Full-Stack Software Developer",
    "location": "Nairobi, Kenya",
    "email": "chepkiruidaphne91@gmail.com",
    "phone": "",
    "phone_display": "",
    "linkedin": "https://www.linkedin.com/in/daphne-chepkirui-382178313",
    "github": ("DMT-Daph", "https://github.com/DMT-Daph"),
    "portfolio": "https://portfolio-omega-umber-28.vercel.app/",
}

SUMMARY = (
    "Full-Stack Software Developer specializing in modern web and mobile applications, "
    "with experience building production systems using React, Next.js, NestJS, Django, React Native, and PostgreSQL. "
    "Experienced in developing enterprise dashboards, APIs, mobile applications, and end-to-end products, "
    "including founding and building BeautiLink. "
    "Leverages AI-assisted development for problem-solving, prototyping, debugging, and delivery "
    "while maintaining engineering and code-quality standards."
)

SKILLS = {
    "Frontend": "JavaScript, TypeScript, React, Next.js, HTML5, CSS3, MUI, Tailwind CSS, Redux",
    "Backend": "Node.js, NestJS, Python, Django, REST APIs",
    "Mobile": "Expo, React Native",
    "Databases": "PostgreSQL, Prisma",
    "Tools and testing": "Git, GitHub, Docker, CI/CD, Jest, React Testing Library, Cypress, Figma",
    "AI and development": "AI-assisted development, prompt engineering",
}

EXPERIENCE = [
    {
        "role": "Software Engineer",
        "company": "Dynamic Mobility Technology",
        "period": "February 2024 - Present",
        "location": "Kenya",
        "bullets": [
            "Delivered KISRS, a healthcare sample referral platform for hospitals and laboratories, implementing RBAC, referral workflows, and dashboard interfaces in Next.js.",
            "Founded and developed BeautiLink, a beauty and wellness marketplace spanning a Next.js admin portal, NestJS APIs, and Expo/React Native apps for booking, shop, chat, and business operations.",
            "Built insurance operations software, including an admin console for agents, claims, commissions, and reporting, and an agent workspace for bookings and customers.",
            "Developed e-Sahal wallet flows on Expo React Native, covering authentication, send money, and transaction history.",
            "Built CarbonFlow reporting and feasibility workflows for carbon project teams, including dashboards and field data screens.",
        ],
    },
    {
        "role": "Freelance Software Engineer",
        "company": "Upeo",
        "period": "2026",
        "location": "Kenya",
        "bullets": [
            "Built the Upeo CBC student portal in Next.js, covering curriculum content, progress tracking, enrollment, and API-backed learner screens.",
        ],
    },
    {
        "role": "Software Engineer",
        "company": "Tenzi Limited",
        "period": "May 2024 - January 2025",
        "location": "Kenya",
        "bullets": [
            "Shipped Tenzi-MRP for procurement, production planning, and inventory, and Tenzi-POS with offline-first IndexedDB sync so checkout continued during outages. Both products are discontinued.",
        ],
    },
    {
        "role": "Software Engineer",
        "company": "CityRight Limited",
        "period": "October 2023 - January 2024",
        "location": "Kenya",
        "bullets": [
            "Built accessible web applications and REST API integrations for e-commerce and marketing sites, with performance and maintainability improvements.",
        ],
    },
    {
        "role": "Software Engineer",
        "company": "Computer Engineering Forum",
        "period": "April 2023 - July 2023",
        "location": "Remote",
        "bullets": [
            "Supported e-commerce builds, performance work, and design collaboration.",
        ],
    },
    {
        "role": "Data Analyst",
        "company": "AMREC",
        "period": "May 2023 - August 2023",
        "location": "Kenya",
        "bullets": [
            "Improved data quality and produced reports and charts for internal teams.",
        ],
    },
]

PROJECTS = [
    {
        "title": "BeautiLink (2026)",
        "stack": "Founder. Next.js, NestJS, Expo, Prisma, PostgreSQL.",
        "desc": "Beauty and wellness marketplace across admin, APIs, and mobile (booking, shop, chat).",
        "links": [("Live demo", "https://beautilink-admin-portal.vercel.app/")],
    },
    {
        "title": "KISRS (2025-2026)",
        "stack": "Software engineer. Next.js, TypeScript, MUI, RBAC.",
        "desc": "Healthcare sample referral platform for hospitals and labs, shipped to UAT.",
        "links": [("Live demo", "https://uat.isrs.co.ke/")],
    },
    {
        "title": "Insurance Admin Portal (2025-2026)",
        "stack": "Software engineer. Next.js, React, MUI, Redux, RBAC.",
        "desc": "Enterprise console for agents, claims, commissions, and reporting.",
        "links": [("Live demo", "http://134.209.70.221:3050/")],
    },
    {
        "title": "e-Sahal (2024-2025)",
        "stack": "Software engineer. Expo, React Native, TypeScript.",
        "desc": "Mobile wallet for the Ethiopian market: auth, transfers, and transactions. Private demo on request.",
        "links": [],
    },
    {
        "title": "CarbonFlow (2026)",
        "stack": "Software engineer. Next.js, TypeScript, NestJS.",
        "desc": "Carbon credits reporting, feasibility studies, and field data dashboards.",
        "links": [("Live demo", "http://carbonflow.sublimematrix.co.ke:3131/auth/login")],
    },
    {
        "title": "Upeo CBC Portal (2026)",
        "stack": "Software engineer. Next.js, TypeScript.",
        "desc": "CBC student portal for curriculum, progress, and learner workflows. Private demo on request.",
        "links": [],
    },
    {
        "title": "Habit Flow (2025-2026)",
        "stack": "Personal project. Next.js, TypeScript, REST API.",
        "desc": "Activity tracker for habits, learning-project checklists, and analytics.",
        "links": [("Live demo", "https://didactic-eureka-psi.vercel.app/")],
    },
]

EDUCATION = [
    ("Bachelor of Science in Computer Science", "Laikipia University, graduated 2023"),
]

VOLUNTEER = [
    (
        "Mentor, Technovation (remote), December 2022 - April 2023",
        "Mentored high-school students on product ideas and team projects.",
    ),
    (
        "Mentor, Google Developer Groups, Laikipia University and Kisumu",
        "Helped developers with frontend basics and study sessions.",
    ),
]

def apply_private() -> None:
    """Phone stays in gitignored docs/resume-private.json. Referee numbers are never written to the public resume."""
    data: dict = {}
    private_path = DOCS / "resume-private.json"
    if private_path.exists():
        data = json.loads(private_path.read_text(encoding="utf-8"))
    CONTACT["phone"] = str(data.get("phone") or os.environ.get("RESUME_PHONE", ""))
    CONTACT["phone_display"] = str(
        data.get("phone_display") or os.environ.get("RESUME_PHONE_DISPLAY", "")
    )


apply_private()

URL_RE = re.compile(
    r"(https?://[^\s\)\],]+|(?:uat\.)?isrs\.co\.ke/?|(?:[\w-]+\.)?vercel\.app[^\s\)]*|github\.com/[\w-]+|carbonflow\.sublimematrix\.co\.ke[^\s\)]*)",
    re.IGNORECASE,
)


def normalize_url(raw: str) -> str:
    url = raw.rstrip(".,;)")
    if url.startswith("http"):
        return url
    if "github.com" in url:
        return f"https://{url}"
    return f"https://{url}"


def text_to_html(text: str) -> str:
    """Turn plain text into HTML with clickable links."""
    escaped = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    def repl(match: re.Match[str]) -> str:
        raw = match.group(0)
        href = normalize_url(raw)
        label = raw.rstrip(".,;)")
        return f'<a href="{href}" color="{BLUE_HEX}"><u>{label}</u></a>'

    return URL_RE.sub(repl, escaped)


def email_to_html(email: str) -> str:
    return f'<a href="mailto:{email}" color="{BLUE_HEX}"><u>{email}</u></a>'


def project_links_md(links: list[tuple[str, str]]) -> str:
    if not links:
        return ""
    return " " + " ".join(f"[{label}]({url})" for label, url in links)


def project_links_html(links: list[tuple[str, str]]) -> str:
    if not links:
        return ""
    parts = [
        f'<a href="{url}" color="{BLUE_HEX}"><u>{label}</u></a>'
        for label, url in links
    ]
    return " &nbsp;".join(parts)


def write_markdown(path: Path) -> None:
    gh_label, gh_url = CONTACT["github"]
    li = CONTACT.get("linkedin")
    li_md = f" | [LinkedIn]({li})" if li else ""
    portfolio = CONTACT.get("portfolio")
    port_md = f" | [Portfolio]({portfolio})" if portfolio else ""
    phone_md = f" | {CONTACT['phone_display']}" if CONTACT.get("phone_display") else ""
    lines = [
        f"# {CONTACT['name'].title()}",
        f"**{CONTACT['title']}** | {CONTACT['location']}",
        f"[{CONTACT['email']}](mailto:{CONTACT['email']}){phone_md}{li_md}{port_md} | [GitHub]({gh_url})",
        "",
        "## Professional Summary",
        SUMMARY,
        "",
        "## Technical Skills",
    ]
    for title, items in SKILLS.items():
        lines.append(f"**{title}:** {items}")
    lines.extend(["", "## Professional Experience"])
    for job in EXPERIENCE:
        lines.append(
            f"### {job['role']} | {job['company']} | {job['location']} | {job['period']}"
        )
        for b in job["bullets"]:
            lines.append(f"- {b}")
    lines.extend(["", "## Selected Projects"])
    for project in PROJECTS:
        lines.append(f"### {project['title']}")
        lines.append(f"*{project['stack']}*")
        lines.append(f"{project['desc']}{project_links_md(project['links'])}")
    lines.extend(["", "## Education"])
    for deg, school in EDUCATION:
        lines.append(f"**{deg}** - {school}")
    lines.extend(["", "## Volunteer and Community"])
    for header, detail in VOLUNTEER:
        lines.append(f"**{header}**")
        lines.append(detail)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


class ResumePDF(FPDF):
    def __init__(self) -> None:
        super().__init__()
        self.set_auto_page_break(auto=True, margin=14)
        self.set_margins(14, 14, 14)

    @property
    def content_width(self) -> float:
        return self.w - self.l_margin - self.r_margin

    def section_heading(self, text: str) -> None:
        if self.get_y() > 250:
            self.add_page()
        self.ln(4)
        self.set_font("Helvetica", "B", 12)
        self.set_text_color(*BLUE)
        self.cell(self.content_width, 7, text, new_x="LMARGIN", new_y="NEXT")
        y = self.get_y()
        self.set_draw_color(*BLUE)
        self.set_line_width(0.4)
        self.line(self.l_margin, y, self.w - self.r_margin, y)
        self.ln(5)
        self.set_text_color(*BLACK)

    def write_html_block(self, html: str, size: int = 10) -> None:
        self.set_font("Helvetica", size=size)
        self.write_html(html)

    def body_text(self, text: str, bold: bool = False) -> None:
        prefix = "<b>" if bold else ""
        suffix = "</b>" if bold else ""
        self.write_html_block(f"<p>{prefix}{text_to_html(text)}{suffix}</p>")

    def bullet(self, text: str) -> None:
        self.write_html_block(f"<p style='margin-left:8px'>- {text_to_html(text)}</p>")


def write_pdf(path: Path) -> None:
    pdf = ResumePDF()
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(*BLUE)
    pdf.cell(0, 10, CONTACT["name"], new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(*BLACK)
    pdf.cell(0, 6, f"{CONTACT['title']}", new_x="LMARGIN", new_y="NEXT")

    gh_label, gh_url = CONTACT["github"]
    gh_html = f'<a href="{gh_url}" color="{BLUE_HEX}"><u>GitHub</u></a>'
    li = CONTACT.get("linkedin", "")
    li_html = (
        f' &nbsp;|&nbsp; <a href="{li}" color="{BLUE_HEX}"><u>LinkedIn</u></a>'
        if li
        else ""
    )
    port = CONTACT.get("portfolio", "")
    port_html = (
        f' &nbsp;|&nbsp; <a href="{port}" color="{BLUE_HEX}"><u>Portfolio</u></a>'
        if port
        else ""
    )
    phone_html = (
        f' &nbsp;|&nbsp; <a href="tel:{CONTACT["phone"]}" color="{BLUE_HEX}"><u>{CONTACT["phone_display"]}</u></a>'
        if CONTACT.get("phone")
        else ""
    )
    contact_html = (
        f"<p>{CONTACT['location']} &nbsp;|&nbsp; "
        f"{email_to_html(CONTACT['email'])}"
        f"{phone_html}"
        f"{li_html}{port_html} &nbsp;|&nbsp; {gh_html}</p>"
    )
    pdf.set_font("Helvetica", "", 10)
    pdf.write_html(contact_html)
    pdf.ln(2)

    pdf.section_heading("PROFESSIONAL SUMMARY")
    pdf.body_text(SUMMARY)

    pdf.section_heading("TECHNICAL SKILLS")
    for title, items in SKILLS.items():
        pdf.write_html_block(
            f'<p><font color="{BLACK_HEX}"><b>{title}:</b></font> '
            f'<font color="{BLACK_HEX}">{items}</font></p>'
        )

    pdf.section_heading("PROFESSIONAL EXPERIENCE")
    for job in EXPERIENCE:
        pdf.write_html_block(
            f'<p><b>{job["role"]}</b> | {job["company"]} | {job["location"]} | <i>{job["period"]}</i></p>'
        )
        for bullet in job["bullets"]:
            pdf.bullet(bullet)
        pdf.ln(1)

    pdf.section_heading("SELECTED PROJECTS")
    for project in PROJECTS:
        pdf.write_html_block(f"<p><b>{text_to_html(project['title'])}</b></p>")
        pdf.body_text(project["stack"])
        links = project_links_html(project["links"])
        desc_html = text_to_html(project["desc"])
        if links:
            pdf.write_html_block(f"<p>{desc_html} &nbsp;{links}</p>")
        else:
            pdf.body_text(project["desc"])
        pdf.ln(1)

    pdf.section_heading("EDUCATION")
    for deg, school in EDUCATION:
        pdf.body_text(f"{deg} | {school}")

    pdf.section_heading("VOLUNTEER AND COMMUNITY")
    for header, detail in VOLUNTEER:
        pdf.body_text(header, bold=True)
        pdf.body_text(detail)

    pdf.output(str(path))


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563EB")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_linked_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    last = 0
    for match in URL_RE.finditer(text):
        if match.start() > last:
            p.add_run(text[last : match.start()])
        raw = match.group(0)
        label = raw.rstrip(".,;)")
        add_hyperlink(p, label, normalize_url(raw))
        last = match.end()
    if last < len(text):
        p.add_run(text[last:])
    if last == 0:
        p.add_run(text)


def add_section_heading(doc: Document, heading: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(heading)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(*BLUE)


def write_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(*BLACK)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(CONTACT["name"].title())
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(*BLUE)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(CONTACT["title"])
    r.font.size = Pt(11)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run(f"{CONTACT['location']} | ")
    add_hyperlink(contact, CONTACT["email"], f"mailto:{CONTACT['email']}")
    if CONTACT.get("phone") and CONTACT.get("phone_display"):
        contact.add_run(" | ")
        add_hyperlink(contact, CONTACT["phone_display"], f"tel:{CONTACT['phone']}")
    if CONTACT.get("linkedin"):
        contact.add_run(" | ")
        add_hyperlink(contact, "LinkedIn", CONTACT["linkedin"])
    if CONTACT.get("portfolio"):
        contact.add_run(" | ")
        add_hyperlink(contact, "Portfolio", CONTACT["portfolio"])
    gh_label, gh_url = CONTACT["github"]
    contact.add_run(" | ")
    add_hyperlink(contact, "GitHub", gh_url)

    add_section_heading(doc, "PROFESSIONAL SUMMARY")
    doc.add_paragraph(SUMMARY)

    add_section_heading(doc, "TECHNICAL SKILLS")
    for skill_title, items in SKILLS.items():
        p = doc.add_paragraph()
        r = p.add_run(f"{skill_title}: ")
        r.bold = True
        r.font.color.rgb = RGBColor(*BLACK)
        r2 = p.add_run(items)
        r2.font.color.rgb = RGBColor(*BLACK)

    add_section_heading(doc, "PROFESSIONAL EXPERIENCE")
    for job in EXPERIENCE:
        p = doc.add_paragraph()
        r = p.add_run(
            f"{job['role']} | {job['company']} | {job['location']} | {job['period']}"
        )
        r.bold = True
        for bullet in job["bullets"]:
            add_linked_paragraph(doc, bullet, style="List Bullet")

    add_section_heading(doc, "SELECTED PROJECTS")
    for project in PROJECTS:
        p = doc.add_paragraph()
        p.add_run(project["title"]).bold = True
        doc.add_paragraph(project["stack"])
        line = doc.add_paragraph()
        line.add_run(project["desc"] + (" " if project["links"] else ""))
        for i, (label, url) in enumerate(project["links"]):
            if i:
                line.add_run(" ")
            add_hyperlink(line, label, url)

    add_section_heading(doc, "EDUCATION")
    for deg, school in EDUCATION:
        doc.add_paragraph(f"{deg} | {school}")

    add_section_heading(doc, "VOLUNTEER AND COMMUNITY")
    for header, detail in VOLUNTEER:
        p = doc.add_paragraph()
        p.add_run(header).bold = True
        doc.add_paragraph(detail)

    doc.save(str(path))


def main() -> None:
    DOCS.mkdir(exist_ok=True)

    md_path = DOCS / "daphne-chepkirui-resume.md"
    docx_path = DOCS / "daphne-chepkirui-resume.docx"

    write_markdown(md_path)
    write_docx(docx_path)
    write_pdf(RESUME_PDF)
    PUBLIC_RESUME.parent.mkdir(exist_ok=True)
    shutil.copyfile(RESUME_PDF, PUBLIC_RESUME)

    nxtlabs = DOCS / "Daphne_Chepkirui_NXTLABS_JavaScript_Developer_CV.docx"
    write_docx(nxtlabs)

    print(f"Wrote {md_path}")
    print(f"Wrote {docx_path}")
    print(f"Wrote {RESUME_PDF}")
    print(f"Wrote {PUBLIC_RESUME}")
    print(f"Wrote {nxtlabs}")


if __name__ == "__main__":
    main()
